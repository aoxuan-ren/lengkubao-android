# -*- coding: utf-8 -*-
"""Add product introduction slide to 洞察性能 PPT."""
from __future__ import annotations

import copy
import io
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

PPT_PATH = Path(r"d:\.文件\创新创业\2026小挑\市赛\专家培训展示PPT（洞察性能3）.pptx")
OUTPUT_PATH = PPT_PATH.with_name(PPT_PATH.stem + "_含产品介绍.pptx")
BACKUP_PATH = PPT_PATH.with_name(PPT_PATH.stem + "_backup.pptx")

# Slide dimensions (13.333 x 7.5 in for 16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TITLE_COLOR = RGBColor(0x00, 0x70, 0xC0)
BODY_COLOR = RGBColor(0x33, 0x33, 0x33)
ACCENT_COLOR = RGBColor(0x00, 0x70, 0xC0)
CARD_BG = RGBColor(0xF2, 0xF7, 0xFC)
FOOTER_COLOR = RGBColor(0x55, 0x55, 0x55)


def analyze_slides(prs: Presentation) -> None:
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        imgs = 0
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip().replace("\n", " ")[:80]
                if t:
                    texts.append(t)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                imgs += 1
        joined = " | ".join(texts[:3])
        print(f"{i}: imgs={imgs} | {joined}")


def clone_picture(slide, source_shape, left, top, width=None, height=None):
    """Clone a picture shape from another slide onto target slide."""
    img_blob = source_shape.image.blob
    w = width or source_shape.width
    h = height or source_shape.height
    slide.shapes.add_picture(io.BytesIO(img_blob), left, top, width=w, height=h)


def add_rounded_rect(slide, left, top, width, height, fill_rgb):
    shape = slide.shapes.add_shape(
        5,  # MSO_SHAPE.ROUNDED_RECTANGLE
        left,
        top,
        width,
        height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
    shape.line.width = Pt(1)
    return shape


def set_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    font_size=14,
    bold=False,
    color=BODY_COLOR,
    align=PP_ALIGN.LEFT,
    font_name="微软雅黑",
):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txBox


def add_bullet_textbox(
    slide,
    left,
    top,
    width,
    height,
    title,
    bullets,
    title_size=16,
    bullet_size=12,
):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run()
    r0.text = title
    r0.font.size = Pt(title_size)
    r0.font.bold = True
    r0.font.color.rgb = ACCENT_COLOR
    r0.font.name = "微软雅黑"

    for bullet in bullets:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_before = Pt(4)
        r = p.add_run()
        r.text = f"• {bullet}"
        r.font.size = Pt(bullet_size)
        r.font.color.rgb = BODY_COLOR
        r.font.name = "微软雅黑"

    return txBox


def get_largest_pictures(slide, n=2):
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    pics.sort(key=lambda s: s.width * s.height, reverse=True)
    return pics[:n]


def add_blank_slide(prs: Presentation, bg_source_index: int | None = 7):
    """Add a blank slide, optionally copying background from a source slide."""
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    new_slide = prs.slides.add_slide(blank_layout)

    for sh in list(new_slide.shapes):
        sp = sh.element
        sp.getparent().remove(sp)

    if bg_source_index is not None:
        source = prs.slides[bg_source_index]
        src_bg = source.element.find(qn("p:cSld")).find(qn("p:bg"))
        if src_bg is not None:
            dst_cSld = new_slide.element.find(qn("p:cSld"))
            old_bg = dst_cSld.find(qn("p:bg"))
            if old_bg is not None:
                dst_cSld.remove(old_bg)
            dst_cSld.insert(0, copy.deepcopy(src_bg))

    return new_slide


def move_slide_to_position(prs: Presentation, from_idx: int, to_idx: int):
    """Move slide from from_idx to to_idx (0-based) in presentation order."""
    sldIdLst = prs.slides._sldIdLst
    sld_ids = list(sldIdLst)
    elem = sld_ids[from_idx]
    sldIdLst.remove(elem)
    sldIdLst.insert(to_idx, elem)


def build_product_slide(slide, prs: Presentation):
    """Build product introduction slide content."""
    # Title
    set_textbox(
        slide,
        Inches(0.6),
        Inches(0.35),
        Inches(5),
        Inches(0.6),
        "产品介绍",
        font_size=32,
        bold=True,
        color=TITLE_COLOR,
    )

    # Subtitle
    set_textbox(
        slide,
        Inches(0.6),
        Inches(0.95),
        Inches(10),
        Inches(0.4),
        "探·算·造·测 一体化低空风场测试产品矩阵",
        font_size=16,
        bold=False,
        color=ACCENT_COLOR,
    )

    # Summary
    set_textbox(
        slide,
        Inches(0.6),
        Inches(1.35),
        Inches(12.1),
        Inches(0.55),
        "洞察性能以「模块化开放风洞 + 高精度动捕 + 智能测控软件」为核心，"
        "为无人机研发与适航验证提供可重复、可量化、低成本的抗风性能测试整体解决方案。",
        font_size=13,
        color=BODY_COLOR,
    )

    # Three column cards
    col_w = Inches(3.85)
    col_h = Inches(3.55)
    col_top = Inches(2.05)
    col_gap = Inches(0.25)
    col_lefts = [Inches(0.6), Inches(0.6) + col_w + col_gap, Inches(0.6) + 2 * (col_w + col_gap)]

    hardware_bullets = [
        "256 单元模块化风扇阵列开放风洞，构型可扩展、风场可定制",
        "空间风场测量子系统，实时三维风速场重构与出口品质控制",
        "Nokov 六自由度光学动捕平台，毫秒级轨迹与姿态同步捕捉",
    ]
    software_bullets = [
        "风场构造与平面风阵设计，支持湍流/侧风/阵风等复杂工况",
        "风扇阵列测控软件，阵列转速联动与测试流程自动化",
        "气动数据采集与姿态分析系统，动捕-风场数据耦合与性能评估",
    ]
    service_bullets = [
        "无人机抗风/气动性能测试与复杂场景验证",
        "风洞平台租赁与小型测试设备交付",
        "适航认证支撑与定制化测试方案（已服务多家行业头部企业）",
    ]

    columns = [
        ("硬件平台", hardware_bullets),
        ("测控软件", software_bullets),
        ("测试服务", service_bullets),
    ]

    for left, (title, bullets) in zip(col_lefts, columns):
        add_rounded_rect(slide, left, col_top, col_w, col_h, CARD_BG)
        add_bullet_textbox(
            slide,
            left + Inches(0.15),
            col_top + Inches(0.12),
            col_w - Inches(0.3),
            col_h - Inches(0.2),
            title,
            bullets,
            title_size=16,
            bullet_size=11,
        )

    # Images from slides 8 and 9 (0-based index 7 and 8)
    slide8 = prs.slides[7]
    slide9 = prs.slides[8]
    pics8 = get_largest_pictures(slide8, 2)
    pics9 = get_largest_pictures(slide9, 1)

    img_top = Inches(5.75)
    if pics8:
        clone_picture(
            slide,
            pics8[0],
            Inches(0.6),
            img_top,
            width=Inches(2.8),
            height=Inches(1.35),
        )
    if len(pics8) > 1:
        clone_picture(
            slide,
            pics8[1],
            Inches(3.55),
            img_top,
            width=Inches(2.8),
            height=Inches(1.35),
        )
    if pics9:
        clone_picture(
            slide,
            pics9[0],
            Inches(6.5),
            img_top,
            width=Inches(2.8),
            height=Inches(1.35),
        )

    # Small icons from slide 10 if available
    slide10 = prs.slides[9]
    small_pics = [
        sh
        for sh in slide10.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width < Inches(1.5)
    ]
    small_pics.sort(key=lambda s: s.width * s.height, reverse=True)
    for i, pic in enumerate(small_pics[:3]):
        clone_picture(
            slide,
            pic,
            Inches(9.55) + Inches(i * 1.05),
            img_top + Inches(0.15),
            width=Inches(0.85),
            height=Inches(0.85),
        )

    # Footer tagline
    set_textbox(
        slide,
        Inches(0.6),
        Inches(7.05),
        Inches(12.1),
        Inches(0.35),
        "相较传统固定风洞：占地面积小、构型灵活、风场类型丰富、测试成本更低——填补低空无人机轻量化适航测试市场空白。",
        font_size=11,
        bold=True,
        color=FOOTER_COLOR,
        align=PP_ALIGN.CENTER,
    )

    # IP footer note
    set_textbox(
        slide,
        Inches(0.6),
        Inches(6.55),
        Inches(12.1),
        Inches(0.3),
        "已获软件著作权 6 项、专利 2 项 | 服务客户：无人机整机/零部件厂商、高校科研院所、国防军工研发单位",
        font_size=9,
        color=FOOTER_COLOR,
        align=PP_ALIGN.CENTER,
    )


def main():
    if not PPT_PATH.exists():
        raise FileNotFoundError(PPT_PATH)

    shutil.copy2(PPT_PATH, BACKUP_PATH)
    print(f"Backup saved to: {BACKUP_PATH}")

    prs = Presentation(str(PPT_PATH))
    print("Before:")
    analyze_slides(prs)

    new_slide = add_blank_slide(prs, bg_source_index=7)
    build_product_slide(new_slide, prs)

    # Move new slide (last) to position 20 (0-based index 20, before thank-you slide 21)
    last_idx = len(prs.slides) - 1
    move_slide_to_position(prs, last_idx, 20)

    save_path = OUTPUT_PATH
    try:
        prs.save(str(PPT_PATH))
        save_path = PPT_PATH
    except PermissionError:
        print("Original file locked; saving to output copy instead.")
        prs.save(str(OUTPUT_PATH))
        save_path = OUTPUT_PATH

    print("\nAfter:")
    prs2 = Presentation(str(save_path))
    analyze_slides(prs2)
    print(f"\nSaved: {save_path}")
    print(f"Total slides: {len(prs2.slides)}")


if __name__ == "__main__":
    main()
